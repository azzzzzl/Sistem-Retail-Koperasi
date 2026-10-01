from django.http import JsonResponse, HttpResponse
from django.shortcuts import render
from django.views.decorators.http import require_GET

from apps.authentication.decorators import permission_required_custom
from apps.master_data.services import MasterDataService
from apps.authentication.repositories import UserRepository

from .services import ReportService

service = ReportService()
master = MasterDataService()


def _params(request):
    return request.GET.get("start_date"), request.GET.get("end_date")


def _want_json(request):
    return request.GET.get("format", "").lower() == "json" or request.headers.get("Accept", "").lower().startswith("application/json")


def _render_or_json(request, template, data):
    if _want_json(request):
        return JsonResponse({"success": True, "data": data}, json_dumps_params={"default": str})
    context = {"report": data, "filters": request.GET}
    if template in {"reports/sales.html", "reports/purchases.html"}:
        context["products"] = master.get_products()
        context["suppliers"] = master.get_suppliers()
        users = UserRepository().find_all()
        for user in users:
            user["id"] = str(user.get("_id"))
        context["users"] = users
    elif template == "reports/suppliers.html":
        context["suppliers"] = master.get_suppliers()
    elif template in {"reports/inventory.html"}:
        context["products"] = master.get_products()
    return render(request, template, context)


def _error(request, exc, template):
    if _want_json(request):
        return JsonResponse({"success": False, "message": str(exc)}, status=400 if isinstance(exc, ValueError) else 500)
    return render(request, template, {"error": str(exc), "filters": request.GET, "report": {}} , status=400 if isinstance(exc, ValueError) else 500)


@require_GET
@permission_required_custom("reports")
def index(request):
    try:
        start, end = _params(request)
        data = service.dashboard(5, start, end)
        return _render_or_json(request, "reports/dashboard.html", data)
    except Exception as exc:
        return _error(request, exc, "reports/dashboard.html")


@require_GET
@permission_required_custom("reports")
def dashboard(request):
    return index(request)


@require_GET
@permission_required_custom("reports")
def sales_report(request):
    try:
        start, end = _params(request)
        data = service.sales_report(start, end, request.GET.get("product_id"), request.GET.get("cashier_id"), request.GET.get("member_id"))
        return _render_or_json(request, "reports/sales.html", data)
    except Exception as exc:
        return _error(request, exc, "reports/sales.html")


@require_GET
@permission_required_custom("reports")
def purchase_report(request):
    try:
        start, end = _params(request)
        data = service.purchase_report(start, end, request.GET.get("supplier_id"), request.GET.get("product_id"))
        return _render_or_json(request, "reports/purchases.html", data)
    except Exception as exc:
        return _error(request, exc, "reports/purchases.html")


@require_GET
@permission_required_custom("reports")
def inventory_report(request):
    try:
        threshold = int(request.GET.get("low_stock_threshold", 5))
        start, end = _params(request)
        data = service.inventory_report(threshold, request.GET.get("product_id"), start, end)
        return _render_or_json(request, "reports/inventory.html", data)
    except Exception as exc:
        return _error(request, exc, "reports/inventory.html")


@require_GET
@permission_required_custom("reports")
def supplier_report(request):
    try:
        start, end = _params(request)
        data = service.supplier_report(request.GET.get("supplier_id"), start, end)
        return _render_or_json(request, "reports/suppliers.html", data)
    except Exception as exc:
        return _error(request, exc, "reports/suppliers.html")


@require_GET
@permission_required_custom("reports")
def payable_report(request):
    try:
        start, end = _params(request)
        data = service.payable_report(start, end, request.GET.get("supplier_id"), request.GET.get("outstanding_only", "false").lower() == "true")
        return _render_or_json(request, "reports/payables.html", data)
    except Exception as exc:
        return _error(request, exc, "reports/payables.html")


@require_GET
@permission_required_custom("reports")
def profit_report(request):
    try:
        start, end = _params(request)
        data = service.profit_report(start, end)
        return _render_or_json(request, "reports/profit.html", data)
    except Exception as exc:
        return _error(request, exc, "reports/profit.html")


@require_GET
@permission_required_custom("audit_log")
def audit_logs(request):
    try:
        limit = min(max(int(request.GET.get("limit", 100)), 1), 500)
        start, end = _params(request)
        data = service.audit_logs(start, end, request.GET.get("action"), request.GET.get("module"), request.GET.get("user_id"), limit)
        return _render_or_json(request, "audit/logs.html", data)
    except Exception as exc:
        return _error(request, exc, "audit/logs.html")


@require_GET
@permission_required_custom("reports")
def export_report(request, report_type, file_format):
    start, end = _params(request)
    try:
        if report_type == "sales": data = service.sales_report(start, end, request.GET.get("product_id"), request.GET.get("cashier_id"), request.GET.get("member_id"))["data"]
        elif report_type == "purchases": data = service.purchase_report(start, end, request.GET.get("supplier_id"), request.GET.get("product_id"))["data"]
        elif report_type == "inventory": data = service.inventory_report(int(request.GET.get("low_stock_threshold",5)), request.GET.get("product_id"), start, end)["stock"]
        elif report_type == "payables": data = service.payable_report(start, end, request.GET.get("supplier_id"), request.GET.get("outstanding_only","false").lower()=="true")["data"]
        elif report_type == "suppliers": data = service.supplier_report(request.GET.get("supplier_id"), start, end)["data"]
        elif report_type == "profit":
            profit = service.profit_report(start, end)
            data = [profit.get("summary", {}) | {"costBasis": profit.get("costBasis", {})}]
        else: raise ValueError("Jenis laporan tidak mendukung export.")

        if file_format == "xlsx":
            from openpyxl import Workbook
            from openpyxl.utils import get_column_letter
            wb=Workbook(); ws=wb.active; ws.title=report_type[:31]
            if data:
                # Flatten common nested values into readable cells.
                headers=list(data[0].keys()); ws.append(headers)
                for row in data:
                    values=[]
                    for h in headers:
                        v=row.get(h)
                        if isinstance(v,(dict,list)): v=str(v)
                        values.append(v)
                    ws.append(values)
                for idx in range(1,len(headers)+1): ws.column_dimensions[get_column_letter(idx)].width=min(max(len(str(headers[idx-1]))+2,12),40)
            response=HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            response["Content-Disposition"]=f'attachment; filename="laporan-{report_type}.xlsx"'
            wb.save(response); return response

        if file_format == "pdf":
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet
            response=HttpResponse(content_type="application/pdf")
            response["Content-Disposition"]=f'attachment; filename="laporan-{report_type}.pdf"'
            doc=SimpleDocTemplate(response,pagesize=landscape(A4),leftMargin=24,rightMargin=24,topMargin=24,bottomMargin=24)
            styles=getSampleStyleSheet(); story=[Paragraph(f"Laporan {report_type.title()}",styles["Title"]),Spacer(1,10)]
            headers=list(data[0].keys()) if data else ["Data"]
            table_data=[headers]
            for row in data:
                table_data.append([str(row.get(h,""))[:70] for h in headers])
            table=Table(table_data,repeatRows=1)
            table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#1f2937")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("GRID",(0,0),(-1,-1),.25,colors.grey),("FONTSIZE",(0,0),(-1,-1),7),("VALIGN",(0,0),(-1,-1),"TOP")]))
            story.append(table); doc.build(story); return response
        raise ValueError("Format export harus xlsx atau pdf.")
    except ImportError as exc:
        return HttpResponse(f"Dependency export belum terpasang: {exc}", status=500)
    except Exception as exc:
        template = "reports/profit.html" if report_type == "profit" else f"reports/{report_type}.html"
        return _error(request, exc, template)
